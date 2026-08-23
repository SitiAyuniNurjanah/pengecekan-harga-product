import streamlit as st
import pandas as pd
from io import BytesIO

from utils.parser import parse_product
from utils.master_loader import load_master
from utils.matcher import match_product

# =========================
# Konfigurasi Halaman
# =========================

st.set_page_config(page_title="Order Price Check", layout="wide")

st.title("Pengecekan Pesanan")
st.write("Upload file pesanan dan Master Pricelist.")

# =========================
# Upload File
# =========================

order_file = st.file_uploader("📄 Upload File Pesanan", type=["xlsx"])

master_file = st.file_uploader("📚 Upload Master Pricelist", type=["xlsx"])

master = None

if master_file is not None:
    master = load_master(master_file)

# =========================
# Fungsi Highlight
# =========================
def highlight_status(row):

    # =========================
    # KHUSUS SET KIT TOREN
    # =========================
    if (
        (
            row.get("Brand") == "PENGUIN"
            and row.get("Jenis") == "SET KIT TOREN"
        )
        or pd.isna(row.get("Status"))
        or str(row.get("Status", "")).strip() == ""
    ):
        return ["background-color:#b3d9ff"] * len(row)

    # =========================
    # TIDAK COCOK
    # =========================
    elif row["Status"] == "TIDAK COCOK":
        return ["background-color:#ffb3b3"] * len(row)

    # =========================
    # TIDAK DITEMUKAN
    # =========================
    elif row["Status"] == "TIDAK DITEMUKAN":
        return ["background-color:#fff3a3"] * len(row)

    return [""] * len(row)

# def highlight_status(row):

#     if row["Status"] == "TIDAK COCOK":
#         return ["background-color:#ffb3b3"] * len(row)

#     elif row["Status"] == "TIDAK DITEMUKAN":
#         return ["background-color:#fff3a3"] * len(row)

#     return [""] * len(row)


# =========================
# Export Excel
# =========================


def convert_excel(df):

    output = BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:

        df.to_excel(writer, index=False, sheet_name="Hasil")

        ws = writer.sheets["Hasil"]

        merah = "FFB3B3"
        kuning = "FFF3A3"

        from openpyxl.styles import PatternFill

        fill_merah = PatternFill(fill_type="solid", start_color=merah)

        fill_kuning = PatternFill(fill_type="solid", start_color=kuning)

        status_col = df.columns.get_loc("Status") + 1

        for row in range(2, len(df) + 2):

            status = ws.cell(row=row, column=status_col).value

            if status == "TIDAK COCOK":

                for col in range(1, len(df.columns) + 1):
                    ws.cell(row=row, column=col).fill = fill_merah

            elif status == "TIDAK DITEMUKAN":

                for col in range(1, len(df.columns) + 1):
                    ws.cell(row=row, column=col).fill = fill_kuning

    output.seek(0)

    return output


# =========================
# Proses Pesanan
# =========================

if order_file is not None:

    df_order = pd.read_excel(order_file)

    if "Nama Barang" not in df_order.columns:

        st.error("Kolom 'Nama Barang' tidak ditemukan")

    else:

        parsed = df_order["Nama Barang"].apply(parse_product)

        df_order["Brand"] = parsed.apply(lambda x: x["brand"])
        df_order["Jenis"] = parsed.apply(lambda x: x["jenis"])
        df_order["Type"] = parsed.apply(lambda x: x["type"])
        df_order["Kode"] = parsed.apply(lambda x: x.get("kode", ""))
        df_order["Ukuran"] = parsed.apply(lambda x: x["ukuran"])
        df_order["Warna"] = parsed.apply(lambda x: x["warna"])

        if master is not None:

            hasil = match_product(df_order, master)

            kolom_tampil = [
                "Tanggal",
                "Nomor #",
                "Kode #",
                "Nama Barang",
                "@Harga",
                "Brand",
                "Jenis",
                "Type",
                "Kode",
                "Ukuran",
                "Warna",
                "Harga Utama",
                "Status",
            ]

            kolom_tampil = [col for col in kolom_tampil if col in hasil.columns]

            hasil = hasil[kolom_tampil]

            st.subheader("Hasil Pencocokan")

            st.dataframe(
                hasil.style.apply(highlight_status, axis=1), use_container_width=True
            )

            excel = convert_excel(hasil)

            st.download_button(
                "📥 Download Hasil",
                data=excel,
                file_name="Hasil_Pengecekan.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        else:

            st.warning("Silakan upload Master Pricelist terlebih dahulu.")
